"""Safely synchronize bank organization and user XLSX exports into the app database.

The command is dry-run by default. It deliberately does not read USER_PASSWORD.
Run ``python scripts/sync_bank_identity.py --help`` for the operational interface.
"""

from __future__ import annotations

import argparse
import getpass
import json
import os
import secrets
import sys
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class SourceOrg:
    row: int
    org_id: str
    org_code: str
    org_name: str
    enabled: bool


@dataclass(frozen=True)
class SourceUser:
    row: int
    user_code: str
    login_code: str
    user_name: str
    org_code: str
    enabled: bool


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="同步行内机构、用户 Excel 到问数应用库（默认仅预检）"
    )
    parser.add_argument("--org-file", type=Path, required=True, help="机构 .xlsx 文件")
    parser.add_argument("--user-file", type=Path, required=True, help="用户 .xlsx 文件")
    parser.add_argument("--org-sheet", help="机构工作表名称；默认第一个工作表")
    parser.add_argument("--user-sheet", help="用户工作表名称；默认第一个工作表")
    parser.add_argument("--org-header-row", type=int, default=1)
    parser.add_argument("--user-header-row", type=int, default=1)
    parser.add_argument("--config", type=Path, required=True, help="后端 backend.env/.env")
    parser.add_argument(
        "--backend-dir",
        type=Path,
        default=Path.cwd(),
        help="包含 src/ask_metric 的 backend-next 目录",
    )
    parser.add_argument(
        "--source-system",
        help="必须与 SSO_SOURCE_SYSTEM 一致；省略时读取配置",
    )
    parser.add_argument("--active-user-status", default="A")
    parser.add_argument("--active-org-status", default="A")
    parser.add_argument(
        "--password-mode",
        choices=("random-unrecoverable", "shared-prompt"),
        default="random-unrecoverable",
        help="新用户密码策略；默认生成不对外提供的随机密码",
    )
    parser.add_argument("--apply", action="store_true", help="实际写库；省略时只预检")
    parser.add_argument("--confirm", help="写库确认口令")
    return parser.parse_args()


def text_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_xlsx(
    path: Path, sheet_name: str | None, header_row: int
) -> list[tuple[int, dict[str, Any]]]:
    if path.suffix.lower() != ".xlsx":
        raise SystemExit("仅支持 .xlsx，请先另存为 xlsx")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("缺少 openpyxl，请按使用说明安装离线 wheel") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise SystemExit("指定工作表不存在，请在源文件中核对工作表名称")
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        header_values = next(
            sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
        )
        headers = [text_cell(value).upper() for value in header_values]
        if len(headers) != len(set(filter(None, headers))):
            raise SystemExit("表头存在重复列，请核对源文件")
        rows: list[tuple[int, dict[str, Any]]] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            row = {headers[index]: value for index, value in enumerate(values) if headers[index]}
            if any(text_cell(value) for value in row.values()):
                rows.append((row_number, row))
        return rows
    finally:
        workbook.close()


def require_columns(
    rows: list[tuple[int, dict[str, Any]]], required: set[str], label: str
) -> None:
    columns = set(rows[0][1]) if rows else set()
    missing = sorted(required - columns)
    if missing:
        raise SystemExit(f"{label}缺少列：{', '.join(missing)}")


def identifier(
    row: dict[str, Any], column: str, label: str, row_number: int, errors: list[str]
) -> str:
    raw = row.get(column)
    value = text_cell(raw)
    if not value:
        errors.append(f"{label}第 {row_number} 行 {column} 为空")
    elif not isinstance(raw, str):
        errors.append(
            f"{label}第 {row_number} 行 {column} 不是文本格式，可能已经丢失前导零"
        )
    return value


def build_sources(
    args: argparse.Namespace,
) -> tuple[dict[str, SourceOrg], list[SourceUser], list[str]]:
    errors: list[str] = []
    org_rows = read_xlsx(args.org_file.resolve(), args.org_sheet, args.org_header_row)
    user_rows = read_xlsx(args.user_file.resolve(), args.user_sheet, args.user_header_row)
    require_columns(org_rows, {"ORG_ID", "ORG_CODE", "ORG_NAME"}, "机构 Excel")
    require_columns(
        user_rows,
        {"USER_CODE", "LOGIN_CODE", "USER_NAME", "ORG_ID"},
        "用户 Excel",
    )

    orgs_by_id: dict[str, SourceOrg] = {}
    org_codes: dict[str, int] = {}
    for row_number, row in org_rows:
        org_id = identifier(row, "ORG_ID", "机构 Excel", row_number, errors)
        org_code = identifier(row, "ORG_CODE", "机构 Excel", row_number, errors)
        org_name = text_cell(row.get("ORG_NAME"))
        status = text_cell(row.get("ORG_STS") or args.active_org_status).upper()
        if not org_name:
            errors.append(f"机构 Excel 第 {row_number} 行 ORG_NAME 为空")
        if len(org_code) > 128 or len(org_name) > 255:
            errors.append(f"机构 Excel 第 {row_number} 行字段超过应用库长度")
        if org_id in orgs_by_id:
            errors.append(
                f"机构 Excel ORG_ID 重复：第 {row_number} 行和"
                f"第 {orgs_by_id[org_id].row} 行"
            )
        if org_code in org_codes:
            errors.append(
                f"机构 Excel ORG_CODE 重复：第 {row_number} 行和"
                f"第 {org_codes[org_code]} 行"
            )
        if org_id and org_code and org_name:
            orgs_by_id[org_id] = SourceOrg(
                row=row_number,
                org_id=org_id,
                org_code=org_code,
                org_name=org_name,
                enabled=status == args.active_org_status.upper(),
            )
            org_codes[org_code] = row_number

    users: list[SourceUser] = []
    user_codes: dict[str, int] = {}
    login_codes: dict[str, int] = {}
    for row_number, row in user_rows:
        user_code = identifier(row, "USER_CODE", "用户 Excel", row_number, errors)
        login_code = identifier(row, "LOGIN_CODE", "用户 Excel", row_number, errors)
        org_id = identifier(row, "ORG_ID", "用户 Excel", row_number, errors)
        user_name = text_cell(row.get("USER_NAME"))
        status = text_cell(row.get("USER_STS") or args.active_user_status).upper()
        org = orgs_by_id.get(org_id)
        if not user_name:
            errors.append(f"用户 Excel 第 {row_number} 行 USER_NAME 为空")
        if org is None:
            errors.append(f"用户 Excel 第 {row_number} 行 ORG_ID 无法关联机构")
        if len(login_code) > 64 or len(user_name) > 128 or len(user_code) > 128:
            errors.append(f"用户 Excel 第 {row_number} 行字段超过应用库长度")
        if user_code in user_codes:
            errors.append(
                f"用户 Excel USER_CODE 重复：第 {row_number} 行和"
                f"第 {user_codes[user_code]} 行"
            )
        login_key = login_code.casefold()
        if login_key in login_codes:
            errors.append(
                f"用户 Excel LOGIN_CODE 重复（忽略大小写）：第 {row_number} 行和"
                f"第 {login_codes[login_key]} 行"
            )
        if user_code and login_code and user_name and org is not None:
            users.append(
                SourceUser(
                    row=row_number,
                    user_code=user_code,
                    login_code=login_code,
                    user_name=user_name,
                    org_code=org.org_code,
                    enabled=(status == args.active_user_status.upper() and org.enabled),
                )
            )
            user_codes[user_code] = row_number
            login_codes[login_key] = row_number
    return orgs_by_id, users, errors


def configure_backend(args: argparse.Namespace) -> str:
    backend_dir = args.backend_dir.expanduser().resolve()
    source_dir = backend_dir / "src"
    if not (source_dir / "ask_metric").is_dir():
        raise SystemExit("--backend-dir 不正确，未找到 src/ask_metric")
    sys.path.insert(0, str(source_dir))
    from dotenv import dotenv_values

    config_file = args.config.expanduser().resolve()
    if not config_file.is_file():
        raise SystemExit("配置文件不存在，请核对 --config")
    for name, value in dotenv_values(config_file).items():
        if value is not None:
            os.environ[name] = value
    from ask_metric.core.config import get_settings
    from ask_metric.infrastructure.db.session import reset_database_runtime

    get_settings.cache_clear()
    reset_database_runtime()
    configured_source = get_settings().sso_source_system.strip()
    source_system = (args.source_system or configured_source).strip()
    if args.source_system and source_system != configured_source:
        raise SystemExit(
            "--source-system 与配置 SSO_SOURCE_SYSTEM 不一致"
        )
    return source_system


def password_hash_factory(mode: str) -> Callable[[], str]:
    """Credentials originate from secure randomness or hidden operator input only."""
    from ask_metric.core.security import hash_password

    if mode == "random-unrecoverable":
        return lambda: hash_password(secrets.token_urlsafe(48))
    if mode != "shared-prompt":
        raise SystemExit("不支持的初始密码策略")
    supplied = getpass.getpass("请输入共享初始密码（不会显示）：")
    confirmation = getpass.getpass("请再次输入共享初始密码：")
    if supplied != confirmation:
        raise SystemExit("两次密码输入不一致")
    if len(supplied) < 12:
        raise SystemExit("共享初始密码至少 12 个字符")
    return lambda: hash_password(supplied)


def synchronize(
    args: argparse.Namespace,
    orgs_by_id: dict[str, SourceOrg],
    users: list[SourceUser],
    source_system: str,
    input_errors: list[str],
) -> int:
    from sqlalchemy import select

    from ask_metric.application.integration_identity import external_user_id
    from ask_metric.infrastructure.db.models import AppUser, OrgTerm
    from ask_metric.infrastructure.db.session import get_app_session_factory

    errors = list(input_errors)
    source_orgs = {org.org_code: org for org in orgs_by_id.values()}
    session = get_app_session_factory()()
    try:
        existing_orgs = {row.org_code: row for row in session.scalars(select(OrgTerm))}
        existing_users = {row.id: row for row in session.scalars(select(AppUser))}
        usernames = {row.username.casefold(): row for row in existing_users.values()}
        user_ids: dict[str, str] = {}
        for user in users:
            user_id = external_user_id(source_system, user.user_code)
            user_ids[user.user_code] = user_id
            owner = usernames.get(user.login_code.casefold())
            if owner is not None and owner.id != user_id:
                errors.append(
                    f"用户 Excel 第 {user.row} 行 LOGIN_CODE 已被其他本地用户占用"
                )

        summary = {
            "mode": "apply" if args.apply else "dry-run",
            "source_organizations": len(source_orgs),
            "source_users": len(users),
            "enabled_source_users": sum(user.enabled for user in users),
            "organizations_to_create": sum(code not in existing_orgs for code in source_orgs),
            "organizations_to_update": sum(
                code in existing_orgs
                and (
                    existing_orgs[code].org_name != org.org_name
                    or existing_orgs[code].enabled != org.enabled
                )
                for code, org in source_orgs.items()
            ),
            "users_to_create": sum(
                user_ids[user.user_code] not in existing_users for user in users
            ),
            "users_to_update": sum(user_ids[user.user_code] in existing_users for user in users),
            "validation_errors": len(errors),
        }
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        if errors:
            print("校验失败（最多显示前 100 条）：", file=sys.stderr)
            for error in errors[:100]:
                print(f"- {error}", file=sys.stderr)
            if len(errors) > 100:
                print(f"- 其余 {len(errors) - 100} 条已省略", file=sys.stderr)
            session.rollback()
            return 2
        if not args.apply:
            session.rollback()
            print("dry-run 完成：未修改数据库。")
            return 0
        if args.confirm != "SYNC_BANK_IDENTITY":
            session.rollback()
            raise SystemExit("拒绝写库：--confirm 必须为 SYNC_BANK_IDENTITY")

        next_password_hash = password_hash_factory(args.password_mode)

        for org in source_orgs.values():
            target = existing_orgs.get(org.org_code)
            if target is None:
                target = OrgTerm(
                    org_code=org.org_code,
                    org_name=org.org_name,
                    aliases=[],
                    enabled=org.enabled,
                )
                session.add(target)
            else:
                target.org_name = org.org_name
                target.enabled = org.enabled

        for source in users:
            user_id = user_ids[source.user_code]
            target = existing_users.get(user_id)
            if target is None:
                target = AppUser(
                    id=user_id,
                    username=source.login_code,
                    password_hash=next_password_hash(),
                    display_name=source.user_name,
                    org_code=source.org_code,
                    role_code="USER",
                    enabled=source.enabled,
                    session_version=0,
                )
                session.add(target)
            else:
                security_changed = (
                    target.username != source.login_code
                    or target.org_code != source.org_code
                    or target.enabled != source.enabled
                )
                target.username = source.login_code
                target.display_name = source.user_name
                target.org_code = source.org_code
                target.enabled = source.enabled
                if security_changed:
                    target.session_version += 1
                # Preserve password_hash and role_code, especially break-glass admins.
        session.commit()
        print("同步成功：事务已提交。Excel 中未出现的用户和机构没有删除或禁用。")
        return 0
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


def main() -> None:
    args = parse_args()
    if args.org_header_row < 1 or args.user_header_row < 1:
        raise SystemExit("表头行号必须大于等于 1")
    try:
        orgs, users, errors = build_sources(args)
        source_system = configure_backend(args)
        result = synchronize(args, orgs, users, source_system, errors)
    except Exception:
        raise SystemExit("同步失败：请核对输入文件、配置和数据库可用性；"
                         "未完成的事务已回滚。") from None
    raise SystemExit(result)


if __name__ == "__main__":
    main()
